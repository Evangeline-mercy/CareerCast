"""
CareerCast - Resume Extraction Module

Purpose:
    1. Extract text from PDF and DOCX resumes.
    2. Extract skills from the extracted resume text.
    3. Return structured resume information.

IMPORTANT:
    This module does NOT perform career prediction.
    It does NOT load any ML model.
    It does NOT access old recommendation files.
"""

from pathlib import Path
import re

import PyPDF2
from docx import Document


# ============================================================
# SUPPORTED FILE TYPES
# ============================================================

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
}


# ============================================================
# SKILL VOCABULARY
# ============================================================

SKILL_VOCABULARY = [
    # Programming
    "python",
    "java",
    "c",
    "c++",
    "c#",
    "javascript",
    "typescript",
    "matlab",
    "verilog",
    "vhdl",
    "sql",
    "r",

    # Machine Learning / AI
    "machine learning",
    "deep learning",
    "artificial intelligence",
    "natural language processing",
    "nlp",
    "computer vision",
    "tensorflow",
    "pytorch",
    "keras",
    "scikit-learn",
    "sklearn",
    "xgboost",
    "pandas",
    "numpy",
    "opencv",

    # Data
    "data analysis",
    "data analytics",
    "data science",
    "data visualization",
    "statistics",
    "power bi",
    "tableau",
    "excel",

    # Web
    "html",
    "css",
    "react",
    "node.js",
    "node",
    "express",
    "flask",
    "django",

    # Database
    "mysql",
    "postgresql",
    "mongodb",
    "oracle",
    "sqlite",

    # Cloud
    "aws",
    "azure",
    "google cloud",
    "gcp",
    "cloud computing",

    # DevOps / Tools
    "git",
    "github",
    "docker",
    "kubernetes",
    "linux",

    # Embedded / ECE
    "arduino",
    "esp32",
    "raspberry pi",
    "embedded systems",
    "iot",
    "internet of things",
    "microcontrollers",
    "8051",
    "8086",
    "fpga",
    "vlsi",
    "embedded c",

    # Other technical skills
    "api",
    "rest api",
    "rest",
    "json",
    "xml",
    "gitlab",
]


# ============================================================
# NORMALIZATION
# ============================================================

def normalize_text(text):
    """
    Normalize extracted resume text while preserving
    meaningful technical terms.
    """

    if not isinstance(text, str):
        raise TypeError("text must be a string.")

    # Normalize line endings
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Replace tabs with spaces
    text = text.replace("\t", " ")

    # Collapse excessive spaces
    text = re.sub(r"[ ]{2,}", " ", text)

    # Collapse excessive blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_text_from_pdf(file_path):
    """
    Extract text from a PDF resume.
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"PDF file not found: {file_path}"
        )

    text_parts = []

    with open(file_path, "rb") as file:

        reader = PyPDF2.PdfReader(file)

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):

            text = page.extract_text()

            if text:
                text_parts.append(text)

    text = "\n".join(text_parts)

    text = normalize_text(text)

    if not text:
        raise ValueError(
            "No text could be extracted from the PDF."
        )

    return text


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_text_from_docx(file_path):
    """
    Extract text from a DOCX resume.
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"DOCX file not found: {file_path}"
        )

    document = Document(file_path)

    text_parts = []

    # Paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text_parts.append(
                    " | ".join(row_text)
                )

    text = "\n".join(text_parts)

    text = normalize_text(text)

    if not text:
        raise ValueError(
            "No text could be extracted from the DOCX."
        )

    return text


# ============================================================
# GENERIC FILE EXTRACTION
# ============================================================

def extract_resume_text(file_path):
    """
    Extract resume text based on file extension.

    Supported:
        .pdf
        .docx
    """

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"Resume file not found: {file_path}"
        )

    extension = file_path.suffix.lower()

    if extension == ".pdf":

        return extract_text_from_pdf(
            file_path
        )

    if extension == ".docx":

        return extract_text_from_docx(
            file_path
        )

    raise ValueError(
        f"Unsupported resume format: {extension}. "
        f"Supported formats: PDF, DOCX."
    )


# ============================================================
# SKILL EXTRACTION
# ============================================================

def extract_skills(resume_text):
    """
    Extract technical skills appearing in the resume.

    Matching is case-insensitive.

    Returns:
        A sorted list of unique skills.
    """

    if not isinstance(resume_text, str):
        raise TypeError(
            "resume_text must be a string."
        )

    text = resume_text.lower()

    found_skills = set()

    for skill in SKILL_VOCABULARY:

        # Escape special regex characters
        escaped_skill = re.escape(
            skill.lower()
        )

        # Word-boundary matching
        pattern = (
            r"(?<!\w)"
            + escaped_skill
            + r"(?!\w)"
        )

        if re.search(pattern, text):

            found_skills.add(skill)

    return sorted(
        found_skills,
        key=lambda x: x.lower(),
    )


# ============================================================
# COMPLETE RESUME ANALYSIS
# ============================================================

def analyze_resume(file_path):
    """
    Complete resume extraction pipeline.

    Returns:

        {
            "file_name": "...",
            "file_type": ".pdf",
            "text": "...",
            "skills": [...],
            "skill_count": N,
            "text_length": N
        }
    """

    file_path = Path(file_path)

    text = extract_resume_text(
        file_path
    )

    skills = extract_skills(
        text
    )

    return {
        "file_name": file_path.name,
        "file_type": file_path.suffix.lower(),
        "text": text,
        "skills": skills,
        "skill_count": len(skills),
        "text_length": len(text),
    }


# ============================================================
# STANDALONE TEST
# ============================================================

if __name__ == "__main__":

    print("=" * 70)
    print("CAREERCAST - RESUME EXTRACTOR TEST")
    print("=" * 70)

    test_text = """
    Electronics and Communication Engineering student.

    Skills:
    Python, SQL, Machine Learning, Deep Learning,
    TensorFlow, Pandas, NumPy, C++, MATLAB,
    Arduino, ESP32, IoT, Embedded Systems,
    Git and GitHub.
    """

    skills = extract_skills(test_text)

    print("\nTEST TEXT:")
    print(test_text)

    print("\nEXTRACTED SKILLS:")
    for index, skill in enumerate(
        skills,
        start=1,
    ):
        print(
            f"{index:2d}. {skill}"
        )

    print(
        f"\nTotal skills detected: "
        f"{len(skills)}"
    )

    if not skills:
        raise RuntimeError(
            "Skill extraction test failed."
        )

    print("\nTEST PASSED.")